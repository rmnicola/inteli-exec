#!/usr/bin/env python3
"""Gera os insumos de cada caso a partir das fontes editáveis.

Convenção de nomes em <caso>/fontes/:
  nome.pdf.md         -> insumos/nome.pdf   (documento formatado)
  nome.manuscrito.md  -> insumos/nome.png   (foto de papel escrito à mão)
  nome.xlsx.csv       -> insumos/nome.xlsx  (planilha; separador ';')
  qualquer outro      -> copiado para insumos/ como está (.txt, .md, .docx...)

PASTAS diz onde ficam as fontes e para onde vão os arquivos gerados. As
políticas institucionais são geradas na própria pasta (ela vai inteira para o
NotebookLM); os arquivos do fluxo 2 do caso 2 vão para insumos/.

Nos .pdf.md, a primeira linha pode escolher o papel timbrado:
  <!-- timbre: valeverde -->  (padrão)
  <!-- timbre: rodar -->
  <!-- timbre: nenhum -->
"""
import csv
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import markdown
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent
CHROMIUM = shutil.which("chromium") or shutil.which("google-chrome")
RODAPE = "Material didático Inteli · empresa e pessoas fictícias"

TIMBRES = {
    "valeverde": """<header class="timbre vv"><div class="marca"><span class="logo">◆</span>
      <div><strong>Vale Verde Distribuidora</strong><small>Limpeza · Higiene · Descartáveis</small></div></div>
      <div class="end">Rod. Anhanguera, km 98 · Campinas/SP</div></header>""",
    "rodar": """<header class="timbre rodar"><div class="marca"><span class="logo">▶▶</span>
      <div><strong>RODAR BRASIL LOGÍSTICA S.A.</strong><small>Soluções em transporte rodoviário de cargas</small></div></div>
      <div class="end">Av. das Nações Unidas, 12.000 · São Paulo/SP</div></header>""",
    "nenhum": "",
}

CSS = """
@page { size: A4; margin: 18mm 18mm 20mm; }
body { font-family: 'Alegreya Sans', 'Adwaita Sans', sans-serif; font-size: 10.5pt; line-height: 1.45; color: #1d2320; }
.timbre { display: flex; justify-content: space-between; align-items: flex-end; padding-bottom: 8px; margin-bottom: 18px; }
.timbre .marca { display: flex; gap: 10px; align-items: center; }
.timbre strong { display: block; font-size: 13pt; letter-spacing: .02em; }
.timbre small { color: #5b6660; font-size: 8.5pt; }
.timbre .end { font-size: 8.5pt; color: #5b6660; }
.vv { border-bottom: 2px solid #2f7d4f; } .vv .logo { color: #2f7d4f; font-size: 22pt; }
.rodar { border-bottom: 3px solid #c2410c; } .rodar .logo { color: #c2410c; font-size: 16pt; font-weight: 700; }
h1 { font-size: 16pt; margin: 0 0 4px; } h2 { font-size: 12.5pt; margin: 18px 0 6px; } h3 { font-size: 11pt; margin: 14px 0 4px; }
table { border-collapse: collapse; width: 100%; margin: 8px 0 12px; font-size: 9.5pt; }
th, td { border: 1px solid #c9d1cc; padding: 4px 7px; text-align: left; vertical-align: top; }
th { background: #eef3f0; }
blockquote { margin: 8px 0; padding: 6px 12px; border-left: 3px solid #9fb5a8; background: #f6f8f7; }
.rodape { position: fixed; bottom: -12mm; left: 0; right: 0; font-size: 7.5pt; color: #8a948f; text-align: center; }
"""

MANUSCRITO_CSS = """
body { margin: 0; background: #d9d4c7; }
.folha { width: 820px; min-height: 1100px; margin: 40px auto; padding: 50px 60px 60px 90px; background:
  repeating-linear-gradient(#fbf8ef 0 33px, #a9c4dd 33px 34px); position: relative;
  transform: rotate(-1.2deg); box-shadow: 0 8px 24px rgba(0,0,0,.25);
  font-family: 'Grape Nuts', cursive; font-size: 23px; line-height: 34px; color: #1f2a6b; }
.folha::before { content: ''; position: absolute; left: 70px; top: 0; bottom: 0; width: 2px; background: #e39a9a; }
.folha h1, .folha h2, .folha h3 { font-size: 30px; line-height: 34px; margin: 0; font-weight: normal; text-decoration: underline; }
.folha p, .folha ul, .folha ol { margin: 0; } .folha table { border-collapse: collapse; }
.folha td, .folha th { border: none; padding: 0 18px 0 0; font-weight: normal; text-align: left; }
.folha strong { font-weight: normal; text-decoration: underline; }
"""


def md_html(text):
    return markdown.markdown(text, extensions=["tables", "sane_lists"])


def chromium(args):
    subprocess.run([CHROMIUM, "--headless=new", "--disable-gpu", "--no-sandbox", *args],
                   check=True, capture_output=True)


def build_pdf(src, dst):
    text = src.read_text(encoding="utf-8")
    m = re.match(r"\s*<!--\s*timbre:\s*(\w+)\s*-->\s*", text)
    timbre = TIMBRES[m.group(1) if m else "valeverde"]
    if m:
        text = text[m.end():]
    html = (f"<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'><style>{CSS}</style></head>"
            f"<body>{timbre}{md_html(text)}<div class='rodape'>{RODAPE}</div></body></html>")
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
        f.write(html)
    chromium(["--no-pdf-header-footer", f"--print-to-pdf={dst}", f"file://{f.name}"])


def build_manuscrito(src, dst):
    html = (f"<!doctype html><html><head><meta charset='utf-8'><style>{MANUSCRITO_CSS}</style></head>"
            f"<body><div class='folha'>{md_html(src.read_text(encoding='utf-8'))}</div></body></html>")
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
        f.write(html)
    chromium(["--hide-scrollbars", "--window-size=1000,1260", f"--screenshot={dst}", f"file://{f.name}"])


def build_xlsx(src, dst):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src.name.split(".")[0][:31]
    with src.open(encoding="utf-8") as f:
        for i, row in enumerate(csv.reader(f, delimiter=";")):
            ws.append([_num(c) for c in row])
            if i == 0:
                for c in ws[1]:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor="EEF3F0")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)
    wb.save(dst)


def _num(v):
    try:
        return int(v)
    except ValueError:
        try:
            return float(v.replace(",", "."))
        except ValueError:
            return v


PASTAS = {
    "politicas-institucionais": "politicas-institucionais",
    "caso-2/fluxo-2": "caso-2/fluxo-2/insumos",
}


def gerar(src, insumos):
    name = src.name
    if name.endswith(".pdf.md"):
        dst = insumos / name.removesuffix(".md")
        build_pdf(src, dst)
    elif name.endswith(".manuscrito.md"):
        dst = insumos / (name.removesuffix(".manuscrito.md") + ".png")
        build_manuscrito(src, dst)
    elif name.endswith(".xlsx.csv"):
        dst = insumos / name.removesuffix(".csv")
        build_xlsx(src, dst)
    else:
        dst = insumos / name
        shutil.copy(src, dst)
    print(f"{src.relative_to(ROOT)} -> {dst.relative_to(ROOT)}")


AGENTE = ROOT / "caso-2/fluxo-2/agente/base"


def sincronizar_agente():
    """Copia a base de conhecimento para a pasta de trabalho do agente (Antigravity)."""
    fluxo2 = ROOT / "caso-2/fluxo-2"
    destinos = {
        "politicas": [(f, f.name.removesuffix(".pdf.md") + ".md") for f in (ROOT / "politicas-institucionais/fontes").glob("*.md")],
        "propostas": [(f, f.name) for f in (fluxo2 / "insumos").glob("proposta-*")],
        "dados": [(f, f.name.removesuffix(".xlsx.csv") + ".csv") for f in (fluxo2 / "fontes").glob("*.xlsx.csv")],
        "pesquisa": [(fluxo2 / "fontes/levantamento-reputacao-transportadoras.pdf.md", "relatorio-reputacao-transportadoras.md")],
    }
    for pasta, arquivos in destinos.items():
        d = AGENTE / pasta
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
        for src, nome in arquivos:
            texto_ou_binario = src.read_bytes()
            if nome.endswith(".md"):
                texto_ou_binario = re.sub(rb"^<!-- timbre: \w+ -->\s*", b"", texto_ou_binario)
            (d / nome).write_bytes(texto_ou_binario)
            print(f"{src.relative_to(ROOT)} -> {(d / nome).relative_to(ROOT)}")


def main():
    for pasta, saida in PASTAS.items():
        insumos = ROOT / saida
        insumos.mkdir(exist_ok=True)
        for src in sorted((ROOT / pasta / "fontes").iterdir()):
            gerar(src, insumos)
    sincronizar_agente()

if __name__ == "__main__":
    main()
